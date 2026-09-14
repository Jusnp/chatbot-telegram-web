from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[2]
DOCUMENTS_DIR = BASE_DIR / "data" / "documentos"
GENERATED_KNOWLEDGE_FILE = BASE_DIR / "data" / "conocimiento_documentos.txt"

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".xlsx"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB

DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)


def _safe_filename(filename: str) -> str:
    """Devuelve un nombre seguro y evita rutas como ../../archivo.txt."""
    name = Path(filename).name.strip()
    if not name:
        raise ValueError("El archivo no tiene un nombre válido.")

    # Conserva letras, números, espacios, guion, punto y guion bajo.
    name = re.sub(r"[^A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9._ -]", "_", name)
    return name


def _unique_path(filename: str) -> Path:
    filename = _safe_filename(filename)
    destination = DOCUMENTS_DIR / filename
    if not destination.exists():
        return destination

    stem = destination.stem
    suffix = destination.suffix
    counter = 2
    while True:
        candidate = DOCUMENTS_DIR / f"{stem}_{counter}{suffix}"
        if not candidate.exists():
            return candidate
        counter += 1


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore").strip()

    if suffix == ".docx":
        document = Document(path)
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]

        # También incluye texto de tablas.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))

        return "\n".join(paragraphs).strip()

    if suffix == ".pdf":
        reader = PdfReader(str(path))
        pages = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"[Página {page_number}]\n{text}")
        return "\n\n".join(pages).strip()

    if suffix == ".xlsx":
        # data_only=True intenta leer el valor calculado de las fórmulas
        # cuando Excel lo haya guardado previamente.
        workbook = load_workbook(filename=path, read_only=True, data_only=True)
        sections = []

        try:
            for sheet in workbook.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    values = []
                    for value in row:
                        if value is None:
                            values.append("")
                        else:
                            values.append(str(value).strip())

                    # Ignora filas completamente vacías.
                    if any(values):
                        rows.append(" | ".join(values))

                if rows:
                    sections.append(
                        f"[Hoja: {sheet.title}]\n" + "\n".join(rows)
                    )
        finally:
            workbook.close()

        return "\n\n".join(sections).strip()

    raise ValueError("Formato no permitido. Usa PDF, DOCX, TXT o XLSX.")


def rebuild_generated_knowledge() -> None:
    sections: list[str] = []

    for path in sorted(DOCUMENTS_DIR.iterdir(), key=lambda p: p.name.lower()):
        if not path.is_file() or path.suffix.lower() not in ALLOWED_EXTENSIONS:
            continue

        try:
            text = extract_text(path)
        except Exception as exc:
            sections.append(
                f"===== FUENTE: {path.name} =====\n"
                f"[No se pudo leer este documento: {exc}]"
            )
            continue

        if text:
            sections.append(f"===== FUENTE: {path.name} =====\n{text}")

    GENERATED_KNOWLEDGE_FILE.write_text(
        "\n\n".join(sections),
        encoding="utf-8",
    )


def save_document(filename: str, content: bytes) -> str:
    safe_name = _safe_filename(filename)
    suffix = Path(safe_name).suffix.lower()

    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Formato no permitido. Solo se aceptan PDF, DOCX, TXT y XLSX.")

    if not content:
        raise ValueError("El archivo está vacío.")

    if len(content) > MAX_FILE_SIZE:
        raise ValueError("El archivo supera el límite de 10 MB.")

    destination = _unique_path(safe_name)
    destination.write_bytes(content)

    # Verifica que realmente se pueda leer antes de incorporarlo.
    try:
        extract_text(destination)
    except Exception:
        destination.unlink(missing_ok=True)
        raise

    rebuild_generated_knowledge()
    return destination.name


def delete_document(filename: str) -> None:
    safe_name = _safe_filename(filename)
    path = DOCUMENTS_DIR / safe_name

    # La comparación evita que un nombre manipulado salga del directorio.
    if path.parent.resolve() != DOCUMENTS_DIR.resolve():
        raise ValueError("Nombre de archivo inválido.")

    if not path.exists() or not path.is_file():
        raise FileNotFoundError("El documento no existe.")

    path.unlink()
    rebuild_generated_knowledge()


def list_documents() -> list[dict]:
    documents = []

    for path in sorted(DOCUMENTS_DIR.iterdir(), key=lambda p: p.name.lower()):
        if not path.is_file() or path.suffix.lower() not in ALLOWED_EXTENSIONS:
            continue
        documents.append(
            {
                "name": path.name,
                "size_kb": round(path.stat().st_size / 1024, 1),
                "type": path.suffix.lower().lstrip(".").upper(),
            }
        )

    return documents
